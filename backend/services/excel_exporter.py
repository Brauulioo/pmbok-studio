"""
Exportación de las 4 matrices oficiales a .xlsx empaquetadas en un .zip.

Las plantillas se cargan con openpyxl.load_workbook(); solo se escriben valores
en la hoja destino y en las coordenadas definidas. El resto de hojas
(Instrucciones, Clasificación) y las celdas combinadas no se tocan.

Origen de cada plantilla, en orden de preferencia:
  1. backend/templates/<archivo>          plantillas originales (locales, fuera del repo)
  2. backend/templates_base/<archivo_base> plantillas propias generadas por
     scripts/crear_plantillas.py (incluidas en el repo)
"""

from __future__ import annotations

import io
import os
import re
import unicodedata
import zipfile
from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from models import KPI, EstadoEntregable, KPIMeta, MAX_MIEMBROS_RACI, Project, ValidationReport
from services.pdf_report import build_pdf
from services.pmbok_validator import (
    clasificar_impacto,
    evaluar_kpi,
    formato_estrategia,
    formato_gravedad,
    formato_probabilidad,
    puntaje_riesgo,
    validate_project,
)

BACKEND_DIR = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = Path(os.getenv("PMBOK_TEMPLATES_DIR", BACKEND_DIR / "templates"))
TEMPLATES_BASE_DIR = Path(os.getenv("PMBOK_TEMPLATES_BASE_DIR", BACKEND_DIR / "templates_base"))
DATE_FORMAT = "DD/MM/YYYY"
AUTOR_EXPORTACION = "PMBOK Studio"


@dataclass(frozen=True)
class TemplateSpec:
    archivo: str          # nombre de la plantilla original
    archivo_base: str     # nombre de la plantilla propia en templates_base/
    hoja: str
    salida: str
    encabezados: Optional[tuple[str, ...]] = None  # Fila 1
    fila_inicio: int = 2
    limpiar_hasta: Optional[int] = None            # None = hasta ws.max_row


DOD = TemplateSpec(
    archivo="Plantilla Matriz DoD.xlsx",
    archivo_base="Matriz DoD.xlsx",
    hoja="Planilla",
    salida="Matriz DoD.xlsx",
    encabezados=("Código", "Entregable", "Criterio de Aceptación", "Responsable", "Fecha de Entrega",
                 "Evidencia/Comprobación"),
    limpiar_hasta=11,
)
RACI = TemplateSpec(
    archivo="PMOinformatica Plantilla Matriz RACI.xls.xlsx",
    archivo_base="Matriz RACI.xlsx",
    hoja="Recursos x Actividad",
    salida="Matriz RACI.xlsx",
    fila_inicio=7,
)
RIESGOS = TemplateSpec(
    archivo="Matriz de Riesgo.xlsx",
    archivo_base="Matriz de Riesgo.xlsx",
    hoja="Registro",
    salida="Matriz de Riesgo.xlsx",
    encabezados=("Código", "Riesgo", "Tipo de Riesgo", "Gravedad", "Probabilidad", "Impacto",
                 "Estrategia de Respuesta", "Responsable", "Fecha de Revisión"),
)
KPIS = TemplateSpec(
    archivo="Matriz de KPIs.xlsx",
    archivo_base="Matriz de KPIs.xlsx",
    hoja="Matriz",
    salida="Matriz de KPIs.xlsx",
    encabezados=("Código", "KPI", "Descripción", "Indicador Adelantado o Rezagado", "Método de Medición",
                 "Meta/Valor Esperado", "Responsable", "Frecuencia"),
)

RACI_FILA_EQUIPO = 6
RACI_COL_PRIMER_MIEMBRO = 3  # C
RACI_COL_ULTIMO_MIEMBRO = RACI_COL_PRIMER_MIEMBRO + MAX_MIEMBROS_RACI - 1  # L
RACI_NCOLS = RACI_COL_ULTIMO_MIEMBRO  # A..L

ALL_SPECS = (DOD, RACI, RIESGOS, KPIS)


class TemplateError(RuntimeError):
    """Plantilla ausente o con estructura distinta a la esperada."""


# --------------------------------------------------------------------------- #
# Utilidades de openpyxl
# --------------------------------------------------------------------------- #
def ruta_plantilla(spec: TemplateSpec) -> Optional[Path]:
    """La original si existe; si no, la plantilla propia de templates_base/."""
    for ruta in (TEMPLATES_DIR / spec.archivo, TEMPLATES_BASE_DIR / spec.archivo_base):
        if ruta.is_file():
            return ruta
    return None


def origen_plantillas() -> dict[str, str]:
    """{salida: 'original' | 'base' | 'faltante'} para /api/health."""
    out = {}
    for s in ALL_SPECS:
        ruta = ruta_plantilla(s)
        out[s.salida] = "faltante" if ruta is None else ("original" if ruta.parent == TEMPLATES_DIR else "base")
    return out


def missing_templates() -> list[str]:
    return [s.archivo for s in ALL_SPECS if ruta_plantilla(s) is None]


def _abrir(spec: TemplateSpec) -> tuple[Any, Worksheet]:
    ruta = ruta_plantilla(spec)
    if ruta is None:
        raise TemplateError(
            f"No se encontró la plantilla '{spec.archivo}' en {TEMPLATES_DIR} ni '{spec.archivo_base}' en "
            f"{TEMPLATES_BASE_DIR}. Genera las plantillas base con: python backend/scripts/crear_plantillas.py"
        )
    wb = load_workbook(ruta)
    if spec.hoja not in wb.sheetnames:
        raise TemplateError(f"'{ruta.name}' no contiene la hoja '{spec.hoja}' (hojas: {wb.sheetnames})")
    return wb, wb[spec.hoja]


def _capturar_estilos(ws: Worksheet, fila: int, ncols: int) -> list[dict[str, Any]]:
    """Copia los estilos de la fila base (una entrada por columna)."""
    estilos = []
    for col in range(1, ncols + 1):
        c = ws.cell(row=fila, column=col)
        estilos.append({
            "font": copy(c.font),
            "border": copy(c.border),
            "fill": copy(c.fill),
            "alignment": copy(c.alignment),
            "number_format": c.number_format,
            "protection": copy(c.protection),
        })
    return estilos


def _limpiar(ws: Worksheet, desde: int, hasta: int, col_ini: int, col_fin: int) -> None:
    """Borra valores conservando el estilo de cada celda."""
    for fila in ws.iter_rows(min_row=desde, max_row=hasta, min_col=col_ini, max_col=col_fin):
        for c in fila:
            if not isinstance(c, MergedCell):
                c.value = None


def _escribir_fila(ws: Worksheet, fila: int, valores: Sequence[Any], estilos: list[dict[str, Any]],
                   col_ini: int = 1, altura: Optional[float] = None) -> None:
    for i, valor in enumerate(valores):
        c = ws.cell(row=fila, column=col_ini + i)
        if isinstance(c, MergedCell):
            continue
        e = estilos[col_ini - 1 + i]
        if not c.has_style:  # filas nuevas fuera del rango de ejemplo
            c.font = copy(e["font"])
            c.border = copy(e["border"])
            c.fill = copy(e["fill"])
            c.alignment = copy(e["alignment"])
            c.protection = copy(e["protection"])
            c.number_format = e["number_format"]
        c.value = valor
        if isinstance(valor, date):
            # openpyxl asigna yyyy-mm-dd al escribir fechas; respetar el formato de la plantilla si lo tiene
            c.number_format = e["number_format"] if e["number_format"] != "General" else DATE_FORMAT
    if altura and ws.row_dimensions[fila].height is None:
        ws.row_dimensions[fila].height = altura


def _ajustar_tablas(ws: Worksheet, ultima_fila: int) -> None:
    """Si la hoja tiene Tablas de Excel o autofiltro, extiende su rango a los datos escritos."""
    for tabla in ws.tables.values():
        ini, fin = tabla.ref.split(":")
        col_fin = re.match(r"[A-Z]+", fin).group(0)
        fila_ini = int(re.search(r"\d+", ini).group(0))
        tabla.ref = f"{ini}:{col_fin}{max(ultima_fila, fila_ini + 1)}"
        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = tabla.ref
    if ws.auto_filter.ref:
        ini, fin = ws.auto_filter.ref.split(":")
        col_fin = re.match(r"[A-Z]+", fin).group(0)
        ws.auto_filter.ref = f"{ini}:{col_fin}{ultima_fila}"


def _llenar_tabla(spec: TemplateSpec, filas: list[Sequence[Any]]) -> bytes:
    """Flujo común DoD / Riesgos / KPIs: encabezados en fila 1 y datos desde la fila 2."""
    wb, ws = _abrir(spec)
    ncols = len(spec.encabezados)
    estilos = _capturar_estilos(ws, spec.fila_inicio, ncols)
    altura = ws.row_dimensions[spec.fila_inicio].height

    for col, titulo in enumerate(spec.encabezados, start=1):
        ws.cell(row=1, column=col).value = titulo

    hasta = spec.limpiar_hasta or ws.max_row
    _limpiar(ws, spec.fila_inicio, max(hasta, spec.fila_inicio), 1, ncols)

    for i, valores in enumerate(filas):
        _escribir_fila(ws, spec.fila_inicio + i, valores, estilos, altura=altura)

    _ajustar_tablas(ws, spec.fila_inicio + max(len(filas), 1) - 1)
    return _guardar(wb)


def _guardar(wb) -> bytes:
    # Que el archivo exportado no arrastre el nombre de quien editó la plantilla por última vez
    wb.properties.lastModifiedBy = AUTOR_EXPORTACION
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Formateo de valores
# --------------------------------------------------------------------------- #
def _codigo_kpi(codigo: str) -> str:
    """K1 → KPI1 (formato de la plantilla); otros códigos se respetan."""
    m = re.fullmatch(r"K(\d+)", codigo, flags=re.IGNORECASE)
    return f"KPI{m.group(1)}" if m else codigo


def _num(v: float) -> str:
    return f"{v:g}"


def _meta_kpi(k: KPI) -> str:
    meta = k.meta
    if isinstance(meta, str):
        return meta
    assert isinstance(meta, KPIMeta)
    unidad = f" {k.unidad}" if k.unidad else ""
    if meta.tipo == "Minimo" and meta.valor is not None:
        texto = f">= {_num(meta.valor)}{unidad}"
    elif meta.tipo == "Maximo" and meta.valor is not None:
        texto = f"<= {_num(meta.valor)}{unidad}"
    elif meta.tipo == "Minimo_referenciado" and meta.referencia:
        texto = f">= {meta.referencia}"
    elif meta.tipo == "Referencia_inicial" and meta.valor is not None:
        texto = f"{_num(meta.valor)}{unidad} (referencia inicial)"
    else:
        texto = meta.tipo.replace("_", " ")
        if meta.valor is not None:
            texto += f": {_num(meta.valor)}{unidad}"
    if meta.meta_cierre is not None:
        texto += f"; al cierre: {_num(meta.meta_cierre)}{unidad}"
    if meta.revision:
        texto += f". {meta.revision}"
    return texto


def _metodo_kpi(k: KPI) -> str:
    if k.metodo_medicion:
        return k.metodo_medicion
    partes = []
    if k.formula:
        partes.append(f"Fórmula: {k.formula}")
    if k.fuente:
        partes.append(f"Fuente: {k.fuente}")
    return ". ".join(partes)


# --------------------------------------------------------------------------- #
# Matrices
# --------------------------------------------------------------------------- #
def _evidencia_dod(d) -> str:
    """Columna 'Evidencia/Comprobación': lo esperado + lo registrado + la aprobación."""
    partes = [d.evidencia] if d.evidencia else []
    if d.evidencia_registrada:
        enlaces = "; ".join(f"{e.url}{f' ({e.descripcion})' if e.descripcion else ''}" for e in d.evidencia_registrada)
        partes.append(f"Registrada: {enlaces}")
    if d.estado == EstadoEntregable.APROBADO and d.aprobacion:
        partes.append(f"APROBADO por {d.aprobacion.rol_aprobador} (semana {d.aprobacion.semana_aprobacion})")
    else:
        partes.append(f"Estado: {d.estado.value}")
    return "\n".join(partes)


def build_dod(p: Project) -> bytes:
    filas = [
        (d.codigo, d.entregable, d.criterio_aceptacion, d.responsable, d.fecha_entrega, _evidencia_dod(d))
        for d in p.dod
    ]
    return _llenar_tabla(DOD, filas)


def build_riesgos(p: Project) -> bytes:
    filas = []
    for r in p.riesgos:
        filas.append((
            r.codigo,
            r.descripcion,
            r.tipo.value,
            formato_gravedad(r.gravedad),
            formato_probabilidad(r.probabilidad),
            clasificar_impacto(puntaje_riesgo(r.gravedad, r.probabilidad)),
            formato_estrategia(r),
            r.propietario,
            r.fecha_revision,
        ))
    return _llenar_tabla(RIESGOS, filas)


ESTADO_KPI = {"cumple": "CUMPLE", "en_riesgo": "EN RIESGO", "no_cumple": "NO CUMPLE", "sin_meta": "SIN META NUMÉRICA"}


def _meta_con_medicion(p: Project, k: KPI) -> str:
    texto = _meta_kpi(k)
    ev = evaluar_kpi(p, k)
    if ev["estado"] == "sin_datos":
        return texto
    valor = f"{ev['valor']:g}".replace(".", ",")
    return f"{texto}\nActual: {valor} (semana {ev['semana']}) · {ESTADO_KPI[ev['estado']]}"


def build_kpis(p: Project) -> bytes:
    filas = [
        (_codigo_kpi(k.codigo), k.nombre, k.definicion, k.tipo_indicador.value, _metodo_kpi(k),
         _meta_con_medicion(p, k), k.responsable, k.frecuencia)
        for k in p.kpis
    ]
    return _llenar_tabla(KPIS, filas)


def build_raci(p: Project) -> bytes:
    """
    Fila 6 (C..L): integrantes. Desde la fila 7: A = ID, B = entregable DoD,
    C..L = R/A/C/I. Las celdas combinadas A5:B5 y C5:L5 no se modifican.
    """
    if len(p.equipo) > MAX_MIEMBROS_RACI:
        raise TemplateError(f"La plantilla RACI admite como máximo {MAX_MIEMBROS_RACI} integrantes.")

    wb, ws = _abrir(RACI)
    estilos_equipo = _capturar_estilos(ws, RACI_FILA_EQUIPO, RACI_NCOLS)
    estilos = _capturar_estilos(ws, RACI.fila_inicio, RACI_NCOLS)
    altura = ws.row_dimensions[RACI.fila_inicio].height

    # Integrantes (C6:L6)
    _limpiar(ws, RACI_FILA_EQUIPO, RACI_FILA_EQUIPO, RACI_COL_PRIMER_MIEMBRO, RACI_COL_ULTIMO_MIEMBRO)
    _escribir_fila(ws, RACI_FILA_EQUIPO, [m.etiqueta for m in p.equipo], estilos_equipo,
                   col_ini=RACI_COL_PRIMER_MIEMBRO)

    # Actividades (A7:L...) en el orden del DoD; actividades RACI sin DoD van al final
    _limpiar(ws, RACI.fila_inicio, max(ws.max_row, RACI.fila_inicio), 1, RACI_NCOLS)
    nombres = {d.codigo: d.entregable for d in p.dod}
    codigos = [d.codigo for d in p.dod] + [c for c in p.raci if c not in nombres]
    for i, cod in enumerate(codigos):
        asign = p.raci.get(cod, {})
        valores = [i + 1, nombres.get(cod, cod)] + [asign.get(m.rol, "") or None for m in p.equipo]
        _escribir_fila(ws, RACI.fila_inicio + i, valores, estilos, altura=altura)

    _ajustar_tablas(ws, RACI.fila_inicio + max(len(codigos), 1) - 1)
    return _guardar(wb)


# --------------------------------------------------------------------------- #
# ZIP
# --------------------------------------------------------------------------- #
def slugify(texto: str) -> str:
    ascii_ = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "_", ascii_).strip("_") or "proyecto"


def export_project_zip(project: Project, report: Optional[ValidationReport] = None) -> io.BytesIO:
    """Genera las 4 matrices + el informe ejecutivo PDF y retorna un BytesIO (posición 0) con el .zip."""
    faltantes = missing_templates()
    if faltantes:
        raise TemplateError(
            f"Faltan plantillas ({', '.join(faltantes)}). Genera las plantillas base con: "
            "python backend/scripts/crear_plantillas.py"
        )

    archivos = {
        DOD.salida: build_dod(project),
        RACI.salida: build_raci(project),
        RIESGOS.salida: build_riesgos(project),
        KPIS.salida: build_kpis(project),
        "Informe Ejecutivo.pdf": build_pdf(project, report or validate_project(project)),
    }
    prefijo = slugify(project.nombre)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for nombre, contenido in archivos.items():
            zf.writestr(f"{prefijo} - {nombre}", contenido)
    buf.seek(0)
    return buf
