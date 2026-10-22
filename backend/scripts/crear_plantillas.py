"""
Genera las plantillas Excel propias de PMBOK Studio en backend/templates_base/.

Reproducen la estructura que espera el exportador (nombres de hoja, encabezados
en la fila 1, filas 5-6 de la RACI con celdas combinadas, estilos en las filas
de datos), con diseño e instrucciones propios. Los encabezados se toman de
services/excel_exporter.py, así plantilla y exportador no se desalinean.

Uso (desde la raíz del repositorio):
    python backend/scripts/crear_plantillas.py
"""

from __future__ import annotations

import sys
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from models import ESTRATEGIAS_PERMITIDAS, MAX_MIEMBROS_RACI  # noqa: E402
from services.excel_exporter import DOD, KPIS, RACI, RIESGOS, TEMPLATES_BASE_DIR  # noqa: E402
from services.pmbok_validator import (  # noqa: E402
    GRAVEDAD_LABELS, PROBABILIDAD_LABELS, clasificar_impacto, formato_gravedad, formato_probabilidad,
)

# Paleta "plano técnico" (la misma de la app)
INK, MUTED, LINE, SOFT, RED = "1B3A5C", "4E6580", "9FB0C2", "E2E9F1", "C2412D"
HEAT = {"Bajo": "DCEBDD", "Medio": "F6E3B8", "Alto": "F3CFC6"}
FUENTE = "Arial"
FILAS_CON_FORMATO = 30
AUTOR = "PMBOK Studio"

_lado = Side(style="thin", color=LINE)
BORDE = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)
ENCABEZADO = dict(
    font=Font(name=FUENTE, bold=True, color="FFFFFF", size=10),
    fill=PatternFill("solid", fgColor=INK),
    alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
    border=BORDE,
)
SUBENCABEZADO = dict(
    font=Font(name=FUENTE, bold=True, color=INK, size=10),
    fill=PatternFill("solid", fgColor=SOFT),
    alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
    border=BORDE,
)
DATO = dict(
    font=Font(name=FUENTE, size=10, color="1F2933"),
    alignment=Alignment(vertical="top", wrap_text=True),
    border=BORDE,
)
CODIGO = dict(DATO, font=Font(name=FUENTE, size=10, bold=True, color=RED))
CENTRADO = dict(DATO, alignment=Alignment(horizontal="center", vertical="top", wrap_text=True))


def _estilo(celda, estilo: dict) -> None:
    for k, v in estilo.items():
        setattr(celda, k, v)


def _libro(titulo: str) -> Workbook:
    wb = Workbook()
    wb.properties.creator = AUTOR
    wb.properties.lastModifiedBy = AUTOR
    wb.properties.title = titulo
    return wb


def _tabla(ws, encabezados, anchos, formatos=None, centradas=()):
    """Encabezados en la fila 1 y FILAS_CON_FORMATO filas de datos con estilo (desde la fila 2)."""
    formatos = formatos or {}
    for col, (titulo, ancho) in enumerate(zip(encabezados, anchos), start=1):
        _estilo(ws.cell(row=1, column=col, value=titulo), ENCABEZADO)
        ws.column_dimensions[get_column_letter(col)].width = ancho
        for fila in range(2, 2 + FILAS_CON_FORMATO):
            c = ws.cell(row=fila, column=col)
            _estilo(c, CODIGO if col == 1 else CENTRADO if col in centradas else DATO)
            if col in formatos:
                c.number_format = formatos[col]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ultima = get_column_letter(len(encabezados))
    ws.auto_filter.ref = f"A1:{ultima}{1 + FILAS_CON_FORMATO}"


def _lista(ws, rango: str, opciones: list[str], titulo: str) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True)
    dv.error, dv.errorTitle = f"Elige un valor de la lista: {', '.join(opciones)}", titulo
    dv.prompt, dv.promptTitle = ", ".join(opciones), titulo
    ws.add_data_validation(dv)
    dv.add(rango)


def _instrucciones(wb: Workbook, titulo: str, pasos: list[str]) -> None:
    ws = wb.create_sheet("Instrucciones")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 100
    ws["B2"] = titulo
    ws["B2"].font = Font(name=FUENTE, bold=True, size=14, color=INK)
    ws["B3"] = "Plantilla de PMBOK Studio. Las hojas de datos se completan automáticamente al emitir el kit."
    ws["B3"].font = Font(name=FUENTE, size=9, color=MUTED)
    for i, texto in enumerate(pasos, start=1):
        fila = 4 + i
        n = ws.cell(row=fila, column=2, value=f"{i:02d}")
        n.font = Font(name=FUENTE, bold=True, size=10, color=RED)
        n.alignment = Alignment(vertical="top")
        t = ws.cell(row=fila, column=3, value=texto)
        t.font = Font(name=FUENTE, size=10, color="1F2933")
        t.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[fila].height = 30 if len(texto) > 95 else 18


# --------------------------------------------------------------------------- #
# Plantillas
# --------------------------------------------------------------------------- #
def plantilla_dod() -> Workbook:
    wb = _libro("Matriz DoD")
    ws = wb.active
    ws.title = DOD.hoja
    _tabla(ws, DOD.encabezados, [10, 32, 48, 24, 16, 46], formatos={5: "DD/MM/YYYY"}, centradas=(5,))
    _instrucciones(wb, "Matriz DoD: Definición de Terminado", [
        "Cada fila es un entregable con un código único (E1, E2, ...).",
        "El criterio de aceptación debe ser verificable: describe qué se medirá o comprobará para darlo por terminado.",
        "El responsable es el rol que ejecuta el entregable y debe figurar como R en la matriz RACI.",
        "Evidencia/Comprobación describe qué se presentará (enlaces, informes, pruebas). Al emitir el kit se agregan "
        "la evidencia registrada y el estado de aprobación.",
        "Un entregable se aprueba solo cuando el rol con A en la RACI verifica la evidencia contra el criterio.",
    ])
    return wb


def plantilla_raci() -> Workbook:
    wb = _libro("Matriz RACI")
    ws = wb.active
    ws.title = RACI.hoja
    ws["A1"] = "Matriz de asignación de responsabilidades (RACI)"
    ws["A1"].font = Font(name=FUENTE, bold=True, size=14, color=INK)
    ws["A2"] = "R = Responsable de ejecutar · A = Aprobador (uno por fila) · C = Consultado · I = Informado"
    ws["A2"].font = Font(name=FUENTE, size=9, color=MUTED)
    ws["A3"] = "Una fila por entregable del DoD, con exactamente una A y al menos una R."
    ws["A3"].font = Font(name=FUENTE, size=9, color=MUTED)

    ultima = get_column_letter(2 + MAX_MIEMBROS_RACI)  # L
    ws.merge_cells("A5:B5")
    ws.merge_cells(f"C5:{ultima}5")
    for col in range(1, 3 + MAX_MIEMBROS_RACI):
        _estilo(ws.cell(row=5, column=col), ENCABEZADO)
        _estilo(ws.cell(row=6, column=col), SUBENCABEZADO)
    ws["A5"] = "Actividad"
    ws["C5"] = "Roles / Responsabilidades"
    ws["A6"] = "ID"
    ws["B6"] = "Entregable / Actividad"
    for i in range(MAX_MIEMBROS_RACI):
        ws.cell(row=6, column=3 + i, value=f"Colaborador {i + 1}")

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 44
    for i in range(MAX_MIEMBROS_RACI):
        ws.column_dimensions[get_column_letter(3 + i)].width = 15
    ws.row_dimensions[6].height = 36

    letra = dict(CENTRADO, font=Font(name=FUENTE, size=11, bold=True, color=INK))
    for fila in range(RACI.fila_inicio, RACI.fila_inicio + FILAS_CON_FORMATO):
        _estilo(ws.cell(row=fila, column=1), CENTRADO)
        _estilo(ws.cell(row=fila, column=2), DATO)
        for col in range(3, 3 + MAX_MIEMBROS_RACI):
            _estilo(ws.cell(row=fila, column=col), letra)
    ws.freeze_panes = "C7"

    rango = f"C{RACI.fila_inicio}:{ultima}200"
    _lista(ws, rango, ["R", "A", "C", "I"], "Asignación RACI")
    ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=['"A"'],
                                  fill=PatternFill("solid", fgColor=INK), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=['"R"'],
                                  fill=PatternFill("solid", fgColor=SOFT), font=Font(color=INK, bold=True)))
    ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=['"C"'], font=Font(color=MUTED)))
    ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=['"I"'], font=Font(color=LINE)))

    _instrucciones(wb, "Matriz RACI", [
        "Fila 6: nombre de cada rol o integrante del equipo, de la columna C a la L (máximo 10).",
        "Desde la fila 7: una fila por entregable, con el mismo nombre que en la Matriz DoD.",
        "En cada celda escribe R, A, C o I, o déjala vacía si el rol no participa.",
        "Cada fila lleva exactamente una A (quien aprueba y responde por el resultado) y al menos una R (quien ejecuta).",
        "Evita que un mismo rol sea R en más de la mitad de los entregables: es una señal de sobrecarga.",
    ])
    return wb


def plantilla_riesgos() -> Workbook:
    wb = _libro("Matriz de Riesgo")
    ws = wb.active
    ws.title = RIESGOS.hoja
    _tabla(ws, RIESGOS.encabezados, [9, 46, 14, 18, 20, 11, 50, 24, 16],
           formatos={9: "DD/MM/YYYY"}, centradas=(3, 6, 9))
    fin = 1 + FILAS_CON_FORMATO
    _lista(ws, f"C2:C{fin}", ["Amenaza", "Oportunidad"], "Tipo de riesgo")
    _lista(ws, f"D2:D{fin}", [formato_gravedad(v) for v in range(1, 6)], "Gravedad")
    _lista(ws, f"E2:E{fin}", [formato_probabilidad(v) for v in range(1, 6)], "Probabilidad")
    _lista(ws, f"F2:F{fin}", ["Bajo", "Medio", "Alto"], "Impacto")
    for nivel, color in HEAT.items():
        ws.conditional_formatting.add(f"F2:F{fin}", CellIsRule(operator="equal", formula=[f'"{nivel}"'],
                                      fill=PatternFill("solid", fgColor=color), font=Font(color=INK, bold=True)))

    # Hoja de clasificación: matriz 5×5 y estrategias
    cl = wb.create_sheet("Clasificación")
    cl.sheet_view.showGridLines = False
    cl["B2"] = "Clasificación de riesgos: puntaje = gravedad × probabilidad"
    cl["B2"].font = Font(name=FUENTE, bold=True, size=14, color=INK)
    cl.column_dimensions["A"].width = 3
    cl.column_dimensions["B"].width = 24
    for col in range(3, 8):
        cl.column_dimensions[get_column_letter(col)].width = 17
    _estilo(cl["B4"], ENCABEZADO)
    cl["B4"] = "Probabilidad \\ Gravedad"
    for g in range(1, 6):
        _estilo(cl.cell(row=4, column=2 + g, value=formato_gravedad(g)), ENCABEZADO)
    cl.row_dimensions[4].height = 32
    for i, p in enumerate(range(5, 0, -1)):
        fila = 5 + i
        _estilo(cl.cell(row=fila, column=2, value=formato_probabilidad(p)), SUBENCABEZADO)
        cl.row_dimensions[fila].height = 26
        for g in range(1, 6):
            puntaje = g * p
            c = cl.cell(row=fila, column=2 + g, value=puntaje)
            _estilo(c, CENTRADO)
            c.fill = PatternFill("solid", fgColor=HEAT[clasificar_impacto(puntaje)])
            c.font = Font(name=FUENTE, bold=True, size=11, color=INK)
            c.alignment = Alignment(horizontal="center", vertical="center")
    fila = 11
    for nivel, rango in (("Bajo", "1 a 6"), ("Medio", "7 a 12"), ("Alto", "13 a 25")):
        c = cl.cell(row=fila, column=2, value=f"{nivel}: puntaje {rango}")
        c.fill = PatternFill("solid", fgColor=HEAT[nivel])
        c.font = Font(name=FUENTE, bold=True, size=10, color=INK)
        c.border = BORDE
        fila += 1
    fila += 1
    _estilo(cl.cell(row=fila, column=2, value="Tipo"), ENCABEZADO)
    _estilo(cl.cell(row=fila, column=3, value="Estrategias de respuesta"), ENCABEZADO)
    cl.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=7)
    for tipo, estrategias in ESTRATEGIAS_PERMITIDAS.items():
        fila += 1
        _estilo(cl.cell(row=fila, column=2, value=tipo.value), DATO)
        _estilo(cl.cell(row=fila, column=3, value=", ".join(e.value for e in estrategias)), DATO)
        cl.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=7)

    _instrucciones(wb, "Matriz de Riesgo", [
        "Cada fila es un riesgo con código único (R1, R2, ...) redactado como causa y efecto: «Si ocurre…, entonces…».",
        "Tipo de riesgo: Amenaza (efecto negativo) u Oportunidad (efecto positivo).",
        "Gravedad y Probabilidad se valoran de 1 a 5 según las escalas de la hoja Clasificación.",
        "Impacto = gravedad × probabilidad: Bajo hasta 6, Medio de 7 a 12 y Alto desde 13.",
        "Estrategia de Respuesta con el formato «Estrategia: descripción». Amenazas: Evitar, Mitigar, Transferir o "
        "Aceptar. Oportunidades: Explotar, Compartir o Mejorar.",
        "Indica un responsable y la fecha en que se volverá a revisar el riesgo.",
    ])
    return wb


def plantilla_kpis() -> Workbook:
    wb = _libro("Matriz de KPIs")
    ws = wb.active
    ws.title = KPIS.hoja
    _tabla(ws, KPIS.encabezados, [9, 26, 44, 18, 40, 32, 24, 16], centradas=(4,))
    _lista(ws, f"D2:D{1 + FILAS_CON_FORMATO}", ["Adelantado", "Rezagado"], "Tipo de indicador")
    _instrucciones(wb, "Matriz de KPIs", [
        "Cada fila es un indicador con código único (KPI1, KPI2, ...).",
        "Indicador adelantado: anticipa resultados (por ejemplo, historias listas para el próximo sprint). "
        "Rezagado: mide resultados ya ocurridos (por ejemplo, SPI o CPI).",
        "Método de Medición: fórmula y fuente del dato, de forma que otra persona obtenga el mismo valor.",
        "Meta/Valor Esperado: valor objetivo con su unidad (por ejemplo «>= 90 %»). Al emitir el kit se agrega la "
        "última medición registrada y si cumple la meta.",
        "Incluye al menos un indicador adelantado para detectar desvíos antes de que ocurran.",
    ])
    return wb


PLANTILLAS = {
    DOD.archivo_base: plantilla_dod,
    RACI.archivo_base: plantilla_raci,
    RIESGOS.archivo_base: plantilla_riesgos,
    KPIS.archivo_base: plantilla_kpis,
}


def main() -> None:
    TEMPLATES_BASE_DIR.mkdir(parents=True, exist_ok=True)
    for nombre, construir in PLANTILLAS.items():
        ruta = TEMPLATES_BASE_DIR / nombre
        construir().save(ruta)
        print(f"Generada {ruta}")


if __name__ == "__main__":
    main()
